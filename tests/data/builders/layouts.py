# SPDX-License-Identifier: BSD-3-Clause
"""Write the synthetic layout fixtures, and the answer key that says how each should load.

    python tests/data/builders/layouts.py [target]

Target defaults to tests/data/synthetic-layouts, and the script refuses to overwrite it. Each file plants the layout
problems found in published spreadsheets, on records borrowed from tests/data/sample/pdf/west-berkshire-reliefs.csv
with the defects real files carry added (dates beside text, `N/A`, whitespace).

The key is the contract. For each file or sheet: `lines` (CSV), each row's fields as written; `dialect`, the settings
that read it as intended, which a wrangler declares and diagnosis should propose; `names` and `rows`, the table those
settings give, each row padded to the names; and `propose`, what diagnosis should find beyond the dialect. Workbook
cells are given as stored, a date as its ISO text, with `date_cells` naming them. Two files are finished by hand in
LibreOffice: layouts.xlsx saved as layouts.xls, and pdf-source.xlsx exported as pdf-source.pdf, sheet by sheet.
"""

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import TYPE_ERROR
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

HERE = Path(__file__).resolve().parent
TARGET = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else HERE.parent / "synthetic-layouts"

HEADER = ["Propref", "Property Address", "Ret Code", "R/V", "Start Date", "Contact"]
# Start Date as text; in a workbook the first and third are date cells, as Excel leaves a day-first column it could
# partly read (the Stockport case). Contact carries N/A, a lone space, an empty cell, and padding.
RECORDS = [
    ["NN00101700004185", "Unit 5, Orpheus House, Calleva Park, Aldermaston, Reading, RG7 8TA", "SBR3", "13250",
     "04/01/2024", "N/A"],
    ["NN03905740027005", "27-31, Meadow Road, Newbury, RG14 7AH", "SBRE", "10250", "13/04/2024", " "],
    ["NN04105100003115", "At Station Road, Pangbourne, Reading, RG8 7AN", "RH0", "8600", "05/02/2024",
     "Parish Clerk"],
    ["NN0010170000401A", "1, Andromeda House, Calleva Park, Aldermaston, Reading, RG7 8AN", "SBR1", "16250",
     "25/12/2024", ""],
    ["NN0010420000111A", "Aldermaston Parish Hall, The Street, Aldermaston, Reading, Berkshire, RG7 4LW", "MAN",
     "5400", "01/02/2026", "N/A"],
    ["NN00201200001117", "The Bell Inn, Aldworth, Reading, RG8 9SE", "PU15", "24000", "30/01/2026", "  Landlord "],
]
DATE_CELLS = {0: date(2024, 1, 4), 2: date(2024, 2, 5)}
CODES = [["Code", "Description"], ["SBR3", "Small Business Rate Relief"], ["RH0", "Retail, Hospitality and Leisure"],
         ["MAN", "Mandatory Charity Relief"]]
FOOTNOTES = [["Source: business rates system, extracted 15/04/2026"], ["* R/V is the rateable value in pounds"]]
# Each a name a script must still resolve to one field: quotes, brackets, comparison signs, a comma, padding, a line
# break, an empty name, a repeat, one name inside another, and the row key's own name.
HOSTILE = ["Owner's name", 'The "Main" Rate', "Rate [2026]", "R/V > 0", "<none>", "Code, Local", " padded ",
           "line\nbreak", "", "Rate", "Rate", "Rate Relief", "__row__"]
MISSING_LIKE = ["N/A", "NA", "null", "NaN", "#N/A", "n/a", "-", "None"]


def field(value, delimiter=","):
    value = "" if value is None else value
    if any(c in value for c in f'{delimiter}"\n\r') or value != value.strip():
        return '"' + value.replace('"', '""') + '"'
    return value


def csv_text(lines, delimiter=",", eol="\r\n"):
    return eol.join(delimiter.join(field(v, delimiter) for v in line) for line in lines) + eol


def padded(rows, width):
    return [row + [""] * (width - len(row)) for row in rows]


def csv_cases():
    """Each case: its file's lines, the text written (where it is not csv_text of the lines), and its key."""
    width = len(HEADER)
    cases = {}

    lines = [["Business rates reliefs, April 2026"], [], [], HEADER, *RECORDS]
    cases["title-above.csv"] = (lines, None, {
        "problem": "a title and two blank lines above the header",
        "dialect": {"skipRows": 3}, "names": HEADER, "rows": RECORDS,
        "propose": {"skipRows": 3, "why": "the first line with as many fields as the data is line 3"}})

    lines = [HEADER, *RECORDS[:3], [], *RECORDS[3:]]
    cases["blank-line-inside.csv"] = (lines, None, {
        "problem": "a blank line inside the data, kept as a row and numbered",
        "dialect": {}, "names": HEADER, "rows": padded([*RECORDS[:3], [], *RECORDS[3:]], width),
        "propose": {"blank_rows": [3]}})

    records = [r.copy() for r in RECORDS]
    records[1][1] = "27-31, Meadow Road\nNewbury, RG14 7AH"
    lines = [HEADER, *records]
    cases["quoted-line-break.csv"] = (lines, None, {
        "problem": "a quoted cell holding a line break, one record over two lines of the file",
        "dialect": {}, "names": HEADER, "rows": records})

    lines = [HEADER, *RECORDS]
    cases["semicolon.csv"] = (lines, csv_text(lines, delimiter=";"), {
        "problem": "`;` between cells, and commas inside the addresses unquoted",
        "dialect": {"delimiter": ";"}, "names": HEADER, "rows": RECORDS,
        "propose": {"delimiter": ";"}})

    # No cell is quoted, so the addresses lose their commas; one cell opens with a quote that never closes
    records = [[v.replace(",", "") for v in r] for r in RECORDS]
    records[2][1] = '"Old Mill Station Road Pangbourne'
    lines = [HEADER, *records]
    text = "\r\n".join(",".join(line) for line in lines) + "\r\n"
    cases["stray-quote.csv"] = (lines, text, {
        "problem": "a cell opening with a quote that never closes, which swallows the rest of the file if read as a "
                   "quote (1.x needed `quoting=csv.QUOTE_NONE`)",
        "dialect": {"quoteChar": None}, "names": HEADER, "rows": records,
        "propose": {"quoteChar": None, "why": "an opening quote with no close before the end of the file"}})

    lines = [[*line, ""] for line in [HEADER, *RECORDS]]
    cases["trailing-delimiter.csv"] = (lines, None, {
        "problem": "every line ends with a delimiter, an empty last column with an empty header",
        "dialect": {}, "names": HEADER, "rows": RECORDS,
        "propose": {"why": "the empty last column is trimmed with the grid, as fp0 trims it"}})

    # Each record over two lines of the file, as every line has the same fields: the second holds one value
    first = HEADER[:5]
    lines = [first, [HEADER[5], "", "", "", ""]]
    for r in RECORDS:
        lines += [r[:5], [r[5], "", "", "", ""]]
    cases["two-line-records.csv"] = (lines, None, {
        "problem": "each record, header included, over two lines: five fields, then one in the first field",
        "dialect": {"rowsPerRecord": [5, 1]}, "names": HEADER, "rows": RECORDS,
        "propose": {"rowsPerRecord": [5, 1], "why": "every second line has one value, in its first field"},
        "refuse": "a record without its second line, or a second line with values past its first field"})

    rows = [[f"v{i}{j}" for j in range(len(HOSTILE))] for i in range(2)]
    lines = [HOSTILE, *rows]
    cases["hostile-headers.csv"] = (lines, None, {
        "problem": "header names a script must still resolve: every name read exactly, none trimmed or joined",
        "dialect": {}, "names_as_written": HOSTILE, "rows": rows,
        "names": [f"Unnamed: {i}" if n == "" else n for i, n in enumerate(HOSTILE[:-3])] + ["Rate.1", "Rate Relief",
                                                                                          "__row__.1"],
        "renamed": "the empty name by position, as stored schemas name it; the repeat and the row key's own name "
                   "by RowKeyParser.deconflict"})

    rows = [[v, str(i)] for i, v in enumerate(MISSING_LIKE)]
    lines = [["Contact", "n"], *rows]
    cases["missing-like-text.csv"] = (lines, None, {
        "problem": "text pandas reads as missing by default; every one is text, and only the source schema can "
                   "say it means missing", "dialect": {}, "names": ["Contact", "n"], "rows": rows})
    return cases


def stored(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def write_records(ws, top, records=RECORDS, header=HEADER, dates=True):
    """Write a header and records from row `top`, and return the coordinates of the date cells."""
    date_cells = []
    for c, name in enumerate(header, start=1):
        ws.cell(row=top, column=c, value=name)
    for i, record in enumerate(records):
        for c, value in enumerate(record, start=1):
            if dates and c == 5 and i in DATE_CELLS:
                cell = ws.cell(row=top + 1 + i, column=c, value=DATE_CELLS[i])
                cell.number_format = "dd/mm/yyyy"
                date_cells.append(cell.coordinate)
            else:
                ws.cell(row=top + 1 + i, column=c, value=value if value != "" else None)
    return date_cells


def workbook_rows(records=RECORDS):
    """The records as a workbook stores them: two dates as date cells, empty text as an empty cell."""
    rows = []
    for i, record in enumerate(records):
        row = [v if v != "" else None for v in record]
        if i in DATE_CELLS:
            row[4] = stored(DATE_CELLS[i])
        rows.append(row)
    return rows


def layouts_workbook():
    wb = Workbook()
    key = {}

    ws = wb.active
    ws.title = "Titled"
    ws["A1"], ws["A2"] = "West Berkshire Council", "Business rates reliefs, April 2026"
    dates = write_records(ws, top=4)
    key["Titled"] = {"problem": "a title and subtitle, then a blank row, above the header on row 4",
                     "dialect": {"skipRows": 3}, "names": HEADER, "rows": workbook_rows(), "date_cells": dates,
                     "propose": {"skipRows": 3}}

    ws = wb.create_sheet("MergedHeader")
    ws["A1"], ws["B1"], ws["D1"] = "Propref", "Property", "Relief"
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:C1")
    ws.merge_cells("D1:E1")
    for c, name in enumerate(["Address", "Postcode", "Code", "R/V"], start=2):
        ws.cell(row=2, column=c, value=name)
    records = [[r[0], r[1].rsplit(", ", 1)[0], r[1].rsplit(", ", 1)[1], r[2], r[3]] for r in RECORDS]
    for i, record in enumerate(records, start=3):
        for c, value in enumerate(record, start=1):
            ws.cell(row=i, column=c, value=value)
    key["MergedHeader"] = {
        "problem": "a header over two rows: one name merged down both, two merged across two columns each",
        "dialect": {"headerRowCount": 2}, "header_rows": [["Propref", "Property", None, "Relief", None],
                                                          [None, "Address", "Postcode", "Code", "R/V"]],
        "merged": ["A1:A2", "B1:C1", "D1:E1"], "rows": records,
        "names": "supplied by the wrangler, one per column: a merged cell over two may be one title or a heading "
                 "over subheadings, so whyqd never builds a name from it",
        "refuse": "headerRowCount 2 without names: the header rows and merged ranges named, and names asked for",
        "propose": {"headerRowCount": 2, "why": "merged ranges in rows 1 and 2, and data from row 3"}}

    ws = wb.create_sheet("BlankColumns")
    names = HEADER[:2] + [None] + HEADER[2:]
    for c, name in enumerate(names, start=1):
        ws.cell(row=1, column=c, value=name)
    for i, record in enumerate(RECORDS, start=2):
        for c, value in enumerate(record[:2] + [None] + record[2:], start=1):
            ws.cell(row=i, column=c, value=value if value != "" else None)
    border = Border(left=Side(style="thin"), right=Side(style="thin"))
    for i in range(1, len(RECORDS) + 2):
        for c in range(len(names) + 1, len(names) + 4):
            ws.cell(row=i, column=c).border = border
    key["BlankColumns"] = {
        "problem": "an empty column inside the table, and three empty but formatted columns after it",
        "dialect": {}, "names": [*HEADER[:2], "Unnamed: 2", *HEADER[2:]],
        "rows": [r[:2] + [None] + r[2:] for r in workbook_rows()],
        "propose": {"empty_columns": [2], "why": "the trailing columns hold formatting only, and are trimmed"}}

    ws = wb.create_sheet("SubTables")
    write_records(ws, top=1, dates=False)
    for r, row in enumerate(CODES, start=len(RECORDS) + 3):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    foot = len(RECORDS) + len(CODES) + 4
    for r, row in enumerate(FOOTNOTES, start=foot):
        ws.cell(row=r, column=1, value=row[0])
    rows = [*RECORDS, [""] * 6, *padded(CODES, 6), [""] * 6, *padded(FOOTNOTES, 6)]
    key["SubTables"] = {
        "problem": "a second table below a blank row, then footnotes; loaded as given, every row is in the table",
        "dialect": {}, "names": HEADER, "rows": [[v if v != "" else None for v in r] for r in rows],
        "propose": {"table_rows": [0, len(RECORDS) - 1],
                    "second_table": {"header_row": len(RECORDS) + 1, "rows": len(CODES) - 1},
                    "footnotes": [len(RECORDS) + len(CODES) + 2, len(rows) - 1]}}

    ws = wb.create_sheet("Errors")
    write_records(ws, top=1, dates=False)
    errors = {"D3": "#N/A", "D5": "#DIV/0!", "D7": "#REF!"}
    for coordinate, value in errors.items():
        ws[coordinate] = value
        ws[coordinate].data_type = TYPE_ERROR
    rows = [[v if v != "" else None for v in r] for r in RECORDS]
    for coordinate, value in errors.items():
        rows[int(coordinate[1:]) - 2][3] = value
    key["Errors"] = {"problem": "error cells in a number column", "dialect": {}, "names": HEADER, "rows": rows,
                     "error_cells": errors}

    ws = wb.create_sheet("Workings")
    ws.sheet_state = "hidden"
    ws["A1"], ws["A2"] = "Total R/V", "=SUM(Titled!D5:D10)"
    key["Workings"] = {"problem": "a hidden sheet, listed and flagged, never read unasked", "state": "hidden"}
    return wb, key


def single_sheet_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Reliefs"
    for c, name in enumerate(CODES[0], start=1):
        ws.cell(row=1, column=c, value=name)
    for r, row in enumerate(CODES[1:], start=3):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    key = {"problem": "the only sheet, with row 2 blank: 1.x's one-row probe judged it empty and dropped it",
           "dialect": {}, "names": CODES[0], "rows": [[None, None], *CODES[1:]], "propose": {"blank_rows": [0]}}
    return wb, key


def printed(rows):
    """Rows as a PDF prints them: a cell's spaces at either end are not kept."""
    return [[v.strip() for v in row] for row in rows]


def pdf_workbook():
    """Sheets for LibreOffice to export as PDF: the kinds of table the sample's PDFs lack."""
    wb = Workbook()
    key = {}
    thin = Side(style="thin")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)

    def fit(ws):
        # A PDF prints a cell only as wide as its column, so each column fits its longest value
        for i, column in enumerate(ws.columns, start=1):
            longest = max(len(str(c.value)) if c.value is not None else 0 for c in column)
            ws.column_dimensions[get_column_letter(i)].width = max(longest, 10) + 2
        # One page wide, as many long as the rows need, so no table's columns are split across pages
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    def rule(ws):
        fit(ws)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None or cell.coordinate in ws.merged_cells:
                    cell.border = grid

    ws = wb.active
    ws.title = "Ruled"
    write_records(ws, top=1)
    rule(ws)
    ws.column_dimensions["B"].width = 30
    for row in ws.iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True)
    key["Ruled"] = {"problem": "a ruled grid, long addresses wrapped over several lines in their cells",
                    "names": HEADER, "rows": printed(RECORDS),
                    "note": "a date cell exports as its displayed text, dd/mm/yyyy, so the rows are the text",
                    "wrapped": "an address extracts with a line break at each wrap, and keeps it: a wrap and a line "
                               "break in the cell look the same in a PDF, so joining is a declared step, never the "
                               "reader's; the ruling keeps each wrapped cell in its one row"}

    ws = wb.create_sheet("MergedHeader")
    ws["A1"], ws["B1"], ws["D1"] = "Propref", "Property", "Relief"
    ws.merge_cells("A1:A2")
    ws.merge_cells("B1:C1")
    ws.merge_cells("D1:E1")
    for c, name in enumerate(["Address", "Postcode", "Code", "R/V"], start=2):
        ws.cell(row=2, column=c, value=name)
    records = [[r[0], r[1].rsplit(", ", 1)[0], r[1].rsplit(", ", 1)[1], r[2], r[3]] for r in RECORDS]
    for i, record in enumerate(records, start=3):
        for c, value in enumerate(record, start=1):
            ws.cell(row=i, column=c, value=value)
    rule(ws)
    key["MergedHeader"] = {"problem": "a ruled header over two rows with merged cells",
                           "header_rows": [["Propref", "Property", "", "Relief", ""],
                                           ["", "Address", "Postcode", "Code", "R/V"]], "rows": records,
                           "names": "supplied by the wrangler, as layouts.xlsx § MergedHeader"}

    ws = wb.create_sheet("TwoTables")
    write_records(ws, top=1, dates=False)
    for r, row in enumerate(CODES, start=len(RECORDS) + 4):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)
    fit(ws)
    key["TwoTables"] = {"problem": "two tables on one page, unruled, two blank rows apart",
                        "tables": [{"names": HEADER, "rows": printed(RECORDS)},
                                   {"names": CODES[0], "rows": CODES[1:]}]}

    ws = wb.create_sheet("Landscape")
    wide = [*HEADER, "Liable From", "Account Holder", "Ward", "Band"]
    extra = [["01/04/2010", "Orpheus Holdings Ltd", "Aldermaston", "B"],
             ["01/04/2017", "Meadow Road Traders", "Newbury Central", "C"],
             ["12/06/2019", "Pangbourne Station Ltd", "Pangbourne", "A"],
             ["01/04/2023", "Andromeda Estates", "Aldermaston", "B"],
             ["01/04/2005", "Aldermaston Parish Council", "Aldermaston", "A"],
             ["30/09/2016", "The Bell Inn (Aldworth) Ltd", "Basildon", "B"]]
    wide_rows = [r + e for r, e in zip(RECORDS, extra, strict=True)]
    write_records(ws, top=1, records=wide_rows, header=wide, dates=False)
    ws.page_setup.orientation = "landscape"
    rule(ws)
    key["Landscape"] = {"problem": "a wide table on a landscape page", "names": wide,
                        "rows": printed(wide_rows)}

    many = [[f"NN{900000000000 + i:014d}", r[1], r[2], r[3], r[4], r[5]]
            for i, r in enumerate(RECORDS * 25)]
    for title, repeat in (("RepeatedHeader", True), ("HeaderOnce", False)):
        ws = wb.create_sheet(title)
        write_records(ws, top=1, records=many, dates=False)
        if repeat:
            ws.print_title_rows = "1:1"
        rule(ws)
        key[title] = {"problem": f"{len(many)} rows over several pages, the header "
                                 + ("repeated on every page" if repeat else "on the first page only"),
                      "names": HEADER, "rows": printed(many)}

    ws = wb.create_sheet("UnruledWrapped")
    write_records(ws, top=1)
    fit(ws)
    ws.column_dimensions["B"].width = 30
    for row in ws.iter_rows(min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True)
    key["UnruledWrapped"] = {
        "problem": "no ruling, and long addresses wrapped in their cells: read by text position, each wrapped line "
                   "becomes a row of its own and words split across columns",
        "names": HEADER, "rows": printed(RECORDS),
        "report": "loaded by the default reading, the rows that look like continuations (an empty first cell, text "
                  "only where a cell wraps) are shown, so the wrangler can declare how the table is read",
        "declared": "with the table's reading declared (its columns' bounds, which rows continue a record), the rows "
                    "above, each wrap a line break in its cell"}
    return wb, key


def main():
    assert not TARGET.exists(), f"{TARGET} exists; fixtures are never overwritten"
    TARGET.mkdir(parents=True)
    key = {"generated": datetime.now().isoformat(timespec="seconds"), "csv": {}, "workbook": {}}

    for name, (lines, text, entry) in csv_cases().items():
        (TARGET / name).write_bytes((text or csv_text(lines)).encode("utf-8"))
        key["csv"][name] = {"lines": lines, **entry}

    wb, sheets = layouts_workbook()
    wb.save(TARGET / "layouts.xlsx")
    key["workbook"]["layouts.xlsx"] = sheets
    key["workbook"]["layouts.xls"] = "layouts.xlsx saved as .xls in LibreOffice; read to the same key"
    wb, sheet = single_sheet_workbook()
    wb.save(TARGET / "single-sheet.xlsx")
    key["workbook"]["single-sheet.xlsx"] = {"Reliefs": sheet}
    wb, sheets = pdf_workbook()
    wb.save(TARGET / "pdf-source.xlsx")
    key["pdf"] = {"pdf-source.pdf": sheets,
                  "made": "pdf-source.xlsx exported in LibreOffice, every sheet, one PDF; each sheet's table "
                          "read across its pages"}

    (TARGET / "answer-key.json").write_text(json.dumps(key, indent=2, ensure_ascii=False, default=str),
                                            encoding="utf-8")
    print(sorted(p.name for p in TARGET.iterdir()))


if __name__ == "__main__":
    main()
