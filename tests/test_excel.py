# SPDX-License-Identifier: BSD-3-Clause
"""A workbook read cell by cell: every sheet, hidden ones included, each cell its stored value in whyqd's own text.

A cell is read as stored, never as displayed, since a displayed short date follows the machine's locale: the number 12
stored under the format `0000000000`, which Excel shows as `0000000012`, reads as `12` (text stored as `0012` stays
`0012`); a date cell reads as its ISO date beside the text in the same column; an error cell as its code. Each sheet is
its own data source, with its own dialect and names.
"""

import json
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

import whyqd as qd
from whyqd.parsers import CellParser, DataSourceParser, RowKeyParser

DATA = Path(__file__).resolve().parent / "data"
FIXTURE = DATA / "synthetic" / "fixture.xlsx"
CELLS = json.loads((DATA / "synthetic" / "answer-key.json").read_text(encoding="utf-8"))["workbook"]
LAYOUTS = DATA / "synthetic-layouts"
WORKBOOKS = json.loads((LAYOUTS / "answer-key.json").read_text(encoding="utf-8"))["workbook"]
LAYOUT = WORKBOOKS["layouts.xlsx"]
ROW = RowKeyParser.ROW


def derive(
    source: Path, mimetype: str, dialect: list[dict] | None = None, names: dict[str, list[str]] | None = None
) -> list[qd.models.DataSourceModel]:
    # A workbook's names are given per sheet, by its name, since each sheet is its own table
    settings = {k: v for k, v in {"dialect": dialect, "names": names}.items() if v is not None}
    datasource = qd.DataSourceDefinition()
    datasource.derive_model(source=source, mimetype=mimetype, **settings)
    models = datasource.get if isinstance(datasource.get, list) else [datasource.get]
    assert all(m.dialect is not None for m in models)
    return models


def sheet(models: list[qd.models.DataSourceModel], name: str) -> tuple[qd.models.DataSourceModel, pd.DataFrame]:
    (model,) = [m for m in models if m.dialect.sheet == name]
    df = DataSourceParser().get(source=model)
    assert df[ROW].tolist() == list(range(len(df)))
    return model, df


def rows(df: pd.DataFrame, names: list[str]) -> list[list[str | None]]:
    return [[None if pd.isna(c) else c for c in row] for row in df[names].values.tolist()]


def stored(cell: dict) -> str | None:
    """The key's stored value as the Python type a reader gives, in whyqd's text."""
    value = cell["stored"]
    if cell["stored_type"] == "d":
        value = datetime.fromisoformat(value) if " " in value else date.fromisoformat(value)
    rendered = CellParser().render(value=value)
    return None if rendered is None else rendered[1]


class TestCells:
    def test_each_cell_reads_as_stored(self):
        # The `Data` sheet's header is row 1 and its records rows 2 to 5; F6 is a formula saved without a value
        cells = {k: v for k, v in CELLS["Data"]["cells"].items() if k[1:] in {"2", "3", "4", "5"}}
        names = list(dict.fromkeys(v["column"] for v in cells.values()))
        expected = [[stored(cells[f"{chr(65 + i)}{r}"]) for i in range(len(names))] for r in range(2, 6)]
        model, df = sheet(derive(FIXTURE, "XLSX"), "Data")
        assert [c.name for c in model.columns] == names
        assert rows(df, names)[:4] == expected

    def test_error_cells_read_as_their_codes(self):
        case = LAYOUT["Errors"]
        model, df = sheet(derive(LAYOUTS / "layouts.xlsx", "XLSX"), "Errors")
        assert rows(df, case["names"]) == case["rows"]


class TestSheets:
    def test_every_sheet_is_read_hidden_and_empty_included(self):
        # Today a sheet whose first two rows read empty is dropped, and the empty one with it
        models = derive(FIXTURE, "XLSX")
        assert [m.dialect.sheet for m in models] == ["Data", "Titled", "Hidden", "Empty"]
        model, df = sheet(models, "Hidden")
        assert [c.name for c in model.columns] == ["secret"]
        assert rows(df, ["secret"]) == [["kept"]]
        model, df = sheet(models, "Empty")
        assert (model.columns, model.index) == ([], 0)

    @pytest.mark.parametrize("workbook", ["layouts.xlsx", "layouts.xls"])
    @pytest.mark.parametrize("name", ["Titled", "BlankColumns", "SubTables", "Errors", "Workings"])
    def test_each_sheet_reads_to_its_key_under_its_dialect(self, workbook, name):
        if workbook == "layouts.xls" and name == "Workings":
            pytest.skip("LibreOffice cached the formula's value in the .xls, so the cell reads 0")
        dialect = [{"sheet": s, **LAYOUT[s].get("dialect", {})} for s in LAYOUT if s != "MergedHeader"]
        mimetype = "XLS" if workbook.endswith(".xls") else "XLSX"
        model, df = sheet(derive(LAYOUTS / workbook, mimetype, dialect), name)
        assert [c.name for c in model.columns] == LAYOUT[name]["names"]
        expected = [row.copy() for row in LAYOUT[name]["rows"]]
        if workbook == "layouts.xls" and name == "Errors":
            # LibreOffice saved the `#REF!` cell as `#NAME?` in the .xls, as xlrd also reads it
            expected[5][3] = "#NAME?"
        assert rows(df, LAYOUT[name]["names"]) == expected

    def test_the_only_sheet_with_a_blank_second_row_reads(self):
        case = WORKBOOKS["single-sheet.xlsx"]["Reliefs"]
        model, df = sheet(derive(LAYOUTS / "single-sheet.xlsx", "XLSX"), "Reliefs")
        assert [c.name for c in model.columns] == case["names"]
        assert rows(df, case["names"]) == case["rows"]

    def test_a_merged_header_takes_the_names_given_for_its_sheet(self):
        # Another sheet keeps its own names, `Propref` among them, never deconflicted against these
        case = LAYOUT["MergedHeader"]
        names = ["Propref", "Address", "Postcode", "Relief Code", "R/V"]
        dialect = [{"sheet": "MergedHeader", **case["dialect"]}, {"sheet": "Titled", **LAYOUT["Titled"]["dialect"]}]
        models = derive(LAYOUTS / "layouts.xlsx", "XLSX", dialect, {"MergedHeader": names})
        model, df = sheet(models, "MergedHeader")
        assert [c.name for c in model.columns] == names
        assert rows(df, names) == case["rows"]
        model, df = sheet(models, "Titled")
        assert [c.name for c in model.columns] == LAYOUT["Titled"]["names"]

    def test_a_merged_header_without_names_is_refused(self):
        with pytest.raises(ValueError, match=r"(?s)headerRowCount.*\bnames\b"):
            derive(LAYOUTS / "layouts.xlsx", "XLSX", [{"sheet": "MergedHeader", "headerRowCount": 2}])


@pytest.fixture
def references(tmp_path) -> Path:
    """Reference cells as councils store them: Gosport's numbers beside text codes, Eastbourne's in scientific format
    with their digits already lost, Thurrock's zeros lost by the publisher, and a UPRN column with blanks."""
    source = tmp_path / "references.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Reference", "UPRN"])
    for reference, uprn, fmt in [
        (769705120, 37019199, "General"),
        ("N30", None, "General"),
        (26620000000000, 100062000000, "0.00E+00"),
        (38100171142560, None, "General"),
        (20020, 10024142866, "General"),
    ]:
        ws.append([reference, uprn])
        for cell in ws[ws.max_row]:
            cell.number_format = fmt
    wb.save(source)
    return source


def test_a_workbook_dated_from_1904_reads_its_own_dates(tmp_path):
    # ECMA-376 § 18.17.4.1: the same stored number is a date four years and a day later under the 1904 system
    source = tmp_path / "dated-1904.xlsx"
    wb = openpyxl.Workbook()
    wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    ws = wb.active
    ws.append(["Start Date"])
    ws.append([date(2024, 4, 1)])
    ws["A2"].number_format = "dd/mm/yyyy"
    wb.save(source)
    datasource = qd.DataSourceDefinition()
    datasource.derive_model(source=source, mimetype="XLSX")
    model = datasource.get[0] if isinstance(datasource.get, list) else datasource.get
    value = DataSourceParser().get(source=model)["Start Date"].iloc[0]
    assert str(value).startswith("2024-04-01")


def test_a_duration_reads_as_iso_8601(tmp_path):
    # A cell formatted `[h]:mm` stores a length of time; a time of day beside it reads as before
    source = tmp_path / "durations.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Hours", "Opens"])
    ws.append([timedelta(hours=25, minutes=30), time(9, 30)])
    ws["A2"].number_format, ws["B2"].number_format = "[h]:mm", "hh:mm"
    wb.save(source)
    (model,) = derive(source, "XLSX")
    assert rows(DataSourceParser().get(source=model), ["Hours", "Opens"]) == [["PT25H30M", "09:30:00"]]


class TestReferences:
    """Each cell renders as stored: never `769705120.0`, never a digit restored or a zero invented."""

    def test_a_reference_reads_as_its_stored_digits(self, references):
        datasource = qd.DataSourceDefinition()
        datasource.derive_model(source=references, mimetype="XLSX")
        model = datasource.get[0] if isinstance(datasource.get, list) else datasource.get
        df = DataSourceParser().get(source=model)
        assert rows(df, ["Reference", "UPRN"]) == [
            ["769705120", "37019199"],
            ["N30", None],
            ["26620000000000", "100062000000"],
            ["38100171142560", None],
            ["20020", "10024142866"],
        ]

    @pytest.mark.xfail(strict=True, raises=(AssertionError, AttributeError), reason="the profile is written at item 5.2")
    def test_a_mixed_column_records_its_cell_types(self, references):
        (model,) = derive(references, "XLSX")
        profile = {c.name: c.profile for c in model.columns}
        assert profile["Reference"].cellTypes == {"n": 4, "s": 1}
        assert (profile["UPRN"].cellTypes, profile["UPRN"].blank) == ({"n": 3}, 2)
