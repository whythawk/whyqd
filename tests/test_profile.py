# SPDX-License-Identifier: BSD-3-Clause
"""What each column holds, as read, in counts: its cells by type, blank and whitespace-only cells, the
anomalies, and, for an `.xlsx`, its number and date cells by their format code. Profiling never changes a cell.

A column's profile counts the table's rows, below its header, each cell as stored, before any trim. Excel's date cells
beside text in one column are the case that needs it: Stockport's `LiabStart` holds 4,232 date cells, swapped by Excel
under a month-first locale, beside 2,382 text cells it could not read so.
"""

from pathlib import Path

import openpyxl
import pytest

import whyqd as qd

DATA = Path(__file__).resolve().parent / "data"
STOCKPORT = DATA / "sample" / "stockport" / "2026-09-04-stockport-original.XLSX"


def profiles(source: Path, mimetype: str, dialect: dict | None = None) -> dict[str, qd.models.ProfileModel]:
    datasource = qd.DataSourceDefinition()
    datasource.derive_model(source=source, mimetype=mimetype, dialect=dialect)
    model = datasource.get[0] if isinstance(datasource.get, list) else datasource.get
    return {c.name: c.profile for c in model.columns}


@pytest.fixture
def anomalies(tmp_path) -> Path:
    cells = [
        "",
        " ",
        "a b",
        " ",
        "x​",
        "﻿y",
        "c\x07",
        "tab\there",
        "0012",
        "0",
        "12/03/2024",
        "2024-03-12",
        "1 Apr 2024",
        "Apr-24",
        "April",
        "1.5",
    ]
    source = tmp_path / "anomalies.csv"
    source.write_text("Title\nValue\n" + "\n".join(f'"{c}"' for c in cells) + "\n", encoding="utf-8", newline="")
    return source


class TestText:
    def test_each_anomaly_is_counted(self, anomalies):
        profile = profiles(anomalies, "CSV", {"skipRows": 1})["Value"].model_dump()
        assert profile == {
            "cellTypes": {"s": 15},
            "blank": 1,
            "whitespace": 2,
            "hardSpace": 2,
            "zeroWidth": 2,
            "control": 1,
            "dateLike": 4,
            "zeroPadded": 1,
            "numberFormats": {},
        }

    def test_the_header_and_rows_above_it_are_not_counted(self, anomalies):
        profile = profiles(anomalies, "CSV")["Title"]
        assert (profile.cellTypes, profile.blank) == ({"s": 16}, 1)

    def test_cells_are_counted_as_stored_before_any_trim(self, anomalies):
        profile = profiles(anomalies, "CSV", {"skipRows": 1, "trim": True})["Value"]
        assert (profile.whitespace, profile.blank) == (2, 1)

    def test_a_derived_data_source_saves_its_profile(self, anomalies):
        datasource = qd.DataSourceDefinition()
        datasource.derive_model(source=anomalies, mimetype="CSV", dialect={"skipRows": 1})
        written = qd.DataSourceDefinition(source=datasource.get.model_dump(by_alias=True, exclude_none=True)).get
        assert written.columns[0].profile and written.columns[0].profile == datasource.get.columns[0].profile


class TestWorkbook:
    def test_excel_date_cells_beside_text_are_recorded(self):
        profile = profiles(STOCKPORT, "XLSX")
        assert profile["LiabStart"].cellTypes == {"s": 2382, "d": 4232}
        assert profile["LiabStart"].numberFormats == {"mm-dd-yy": 4232}

    def test_number_cells_are_counted_by_format_code(self):
        profile = profiles(STOCKPORT, "XLSX")
        assert profile["PropAddress1"].numberFormats == {"General": 12, "d-mmm": 5}
        assert profile["Address1"].numberFormats == {"General": 201, "0.00E+00": 79}

    def test_a_custom_format_is_read_from_the_workbook(self, tmp_path):
        source = tmp_path / "formats.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Reference"])
        for value, fmt in [(12, "0000000000"), (13, "0000000000"), ("0014", "@"), (15, "General")]:
            ws.append([value])
            ws.cell(row=ws.max_row, column=1).number_format = fmt
        wb.save(source)
        profile = profiles(source, "XLSX")["Reference"]
        assert profile.cellTypes == {"s": 1, "n": 3}
        assert profile.numberFormats == {"0000000000": 2, "General": 1}
        assert profile.zeroPadded == 1

    def test_an_xls_has_no_number_formats(self):
        profile = profiles(DATA / "synthetic-layouts" / "layouts.xls", "XLS", {"sheet": "Titled", "skipRows": 3})
        assert any({"n", "d"} & set(p.cellTypes) for p in profile.values())
        assert all(p.numberFormats == {} for p in profile.values())
