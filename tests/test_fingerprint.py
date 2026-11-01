# SPDX-License-Identifier: BSD-3-Clause
"""Two proofs of a source: its file's byte hash, which any byte changes, and `fp0`, which only its data changes.

`fp0` hashes one sheet's grid as parsed, or a file's without sheets, each cell its type and text, before any row is
skipped, joined, or taken as the header, and before any cell is trimmed: it proves the source, not the selection. Its
bytes are written out here, so its definition is held apart from any reader. A file saved again with a sheet's data
unchanged verifies that sheet with a warning. A PDF has the byte hash only, since its grid depends on how it is read.
"""

import hashlib
from pathlib import Path
import warnings

import openpyxl
import polars as pl
import pytest
import rfc8785

import whyqd as qd
from whyqd.parsers import CoreParser, FingerprintParser

DATA = Path(__file__).resolve().parent / "data"
SYNTHETIC = DATA / "synthetic"
ENCODINGS = {
    "fixture-utf8.csv": None,
    "fixture-utf8-bom.csv": None,
    "fixture-utf8-crlf.csv": None,
    "fixture-utf8-mixed-eol.csv": None,
    "fixture-utf16-bom.csv": None,
    "fixture-latin-1.csv": None,
    "fixture-windows-1252.csv": None,
    "fixture-mac-roman.csv": "macintosh",
}


def derive(source: Path, mimetype: str, dialect: dict | None = None) -> qd.DataSourceDefinition:
    datasource = qd.DataSourceDefinition()
    datasource.derive_model(source=source, mimetype=mimetype, dialect=dialect)
    return datasource


def workbook(path: Path, cells: dict[str, object], sheets: tuple[str, ...] = ("Data",)) -> Path:
    wb = openpyxl.Workbook()
    wb.active.title = sheets[0]
    for name in sheets[1:]:
        wb.create_sheet(name)
    for ws in wb.worksheets:
        for ref, value in cells.items():
            ws[ref] = value
    wb.save(path)
    return path


class TestFp0:
    def test_one_file_in_any_encoding_has_one_fp0(self):
        models = [
            derive(SYNTHETIC / name, "CSV", None if encoding is None else {"encoding": encoding}).get
            for name, encoding in ENCODINGS.items()
        ]
        assert len({m.get_checksum("fp0") for m in models}) == 1
        # The Latin-1 and Windows-1252 fixtures hold no byte that tells the two apart, so are one file
        assert len({m.get_checksum("bytes") for m in models}) == len({(SYNTHETIC / n).read_bytes() for n in ENCODINGS})

    def test_fp0_is_blake2b_512(self):
        model = derive(SYNTHETIC / "fixture-utf8.csv", "CSV").get
        fp0 = next(c for c in model.checksums if c.content == "fp0")
        assert fp0.algorithm == "BLAKE2b-512"
        assert len(fp0.checksumValue) == 128

    def test_fp0_is_taken_before_the_dialect_selects(self, tmp_path):
        source = tmp_path / "padded.csv"
        source.write_text("Title\nName,Code\n a ,1\n", encoding="utf-8")
        plain = derive(source, "CSV").get
        selected = derive(source, "CSV", {"skipRows": 1, "trim": True}).get
        assert [c.name for c in selected.columns] == ["Name", "Code"]
        assert plain.get_checksum("fp0") and selected.get_checksum("fp0") == plain.get_checksum("fp0")

    def test_a_trimmed_cell_reads_trimmed(self, tmp_path):
        source = tmp_path / "padded.csv"
        source.write_text("Name,Code\n a ,1\n", encoding="utf-8")
        datasource = derive(source, "CSV", {"trim": True})
        df = datasource.get_data()
        assert df["Name"].tolist() == ["a"]

    def test_each_sheet_has_its_own_fp0(self, tmp_path):
        source = workbook(tmp_path / "two.xlsx", {"A1": "Name", "A2": "a"}, sheets=("One", "Two"))
        wb = openpyxl.load_workbook(source)
        wb["Two"]["A2"] = "b"
        wb.save(source)
        one, two = derive(source, "XLSX").get
        assert one.get_checksum("fp0") != two.get_checksum("fp0")

    def test_a_sheets_fp0_is_its_alone(self, tmp_path):
        alone = derive(workbook(tmp_path / "one.xlsx", {"A1": "Name", "A2": "a"}), "XLSX").get
        among = derive(workbook(tmp_path / "two.xlsx", {"A1": "Name", "A2": "a"}, sheets=("Data", "More")), "XLSX").get
        assert alone.get_checksum("fp0") == among[0].get_checksum("fp0")

    def test_a_sheets_name_is_not_in_fp0(self, tmp_path):
        one = derive(workbook(tmp_path / "one.xlsx", {"A1": "Name"}, sheets=("One",)), "XLSX").get
        two = derive(workbook(tmp_path / "two.xlsx", {"A1": "Name"}, sheets=("Two",)), "XLSX").get
        assert one.get_checksum("fp0") == two.get_checksum("fp0")

    def test_a_cells_stored_type_is_in_fp0(self, tmp_path):
        number = derive(workbook(tmp_path / "number.xlsx", {"A1": "Code", "A2": 12}), "XLSX").get
        text = derive(workbook(tmp_path / "text.xlsx", {"A1": "Code", "A2": "12"}), "XLSX").get
        assert number.get_checksum("fp0") != text.get_checksum("fp0")

    def test_a_pdf_has_the_byte_hash_only(self):
        pytest.importorskip("pdfplumber")
        model = derive(DATA / "sample" / "pdf" / "NNDR EXEMPTIONS.pdf", "PDF").get
        assert model.get_checksum("bytes")
        assert model.get_checksum("fp0") is None


class TestForm:
    """The bytes written by hand and hashed with `hashlib` alone, so a change to the form, or to any reader, fails."""

    PLAIN = b'fp0\n[["s","Code"],["s","Name"]]\n[["s","12"],["s","a"]]'
    PLAIN_DIGEST = (
        "a4274d48092cbecf450d8fbd77d01e8ba74188232e35bacabf6bb5a02e4cc08d"
        "fb405a716fdbe336550046ff82cd3d61769cd1386fa1117ab598310030876f13"
    )
    AWKWARD_TEXT = 'a"\\\n \U0001f600\x01\x7fé'
    AWKWARD = 'fp0\n[["s","Code"],null]\n[["n","12"],["s","a\\"\\\\\\n \U0001f600\\u0001\x7fé"]]\n[null,null]'
    AWKWARD_DIGEST = (
        "b78278d120d0adf816b251b02f83f6cf3c12b9c5c9fc0c17e4c20983c276f6c0"
        "e4508756e15958b74e01b036d32e606cc2f757c2308a488704b0704eff38984f"
    )

    def awkward(self) -> tuple[pl.DataFrame, pl.DataFrame]:
        grid = pl.DataFrame({"column_0": ["Code", "12", None], "column_1": [None, self.AWKWARD_TEXT, None]})
        types = pl.DataFrame({"column_0": ["s", "n", None], "column_1": [None, "s", None]})
        return grid, types

    def test_the_digests_are_of_the_bytes_written_here(self):
        assert hashlib.blake2b(self.PLAIN).hexdigest() == self.PLAIN_DIGEST
        assert hashlib.blake2b(self.AWKWARD.encode("utf-8")).hexdigest() == self.AWKWARD_DIGEST

    def test_fp0_hashes_those_bytes(self):
        grid, types = self.awkward()
        assert FingerprintParser().fp0(grid=grid, types=types) == self.AWKWARD_DIGEST

    def test_those_bytes_are_rfc_8785(self):
        grid, types = self.awkward()
        lines = [
            rfc8785.dumps([None if text is None else [tag, text] for text, tag in zip(texts, tags, strict=True)])
            for texts, tags in zip(grid.iter_rows(), types.iter_rows(), strict=True)
        ]
        assert b"\n".join([b"fp0", *lines]) == self.AWKWARD.encode("utf-8")

    def test_a_csv_reaches_the_same_digest(self, tmp_path):
        source = tmp_path / "plain.csv"
        source.write_text("Code,Name\n12,a\n", encoding="utf-8")
        assert derive(source, "CSV").get.get_checksum("fp0") == self.PLAIN_DIGEST

    def test_a_workbook_reaches_the_same_digest(self, tmp_path):
        source = workbook(tmp_path / "plain.xlsx", {"A1": "Code", "B1": "Name", "A2": "12", "B2": "a"})
        assert derive(source, "XLSX").get.get_checksum("fp0") == self.PLAIN_DIGEST


class TestVerify:
    def test_another_sheet_changing_leaves_a_sheet_saved_again(self, tmp_path):
        source = workbook(tmp_path / "two.xlsx", {"A1": "Name", "A2": "a"}, sheets=("One", "Two"))
        one, two = (qd.DataSourceDefinition(source=m) for m in derive(source, "XLSX").get)
        wb = openpyxl.load_workbook(source)
        wb["Two"]["A2"] = "b"
        wb.save(source)
        with pytest.warns(UserWarning, match="saved again"):
            assert one.validate()
        with pytest.raises(AssertionError, match="Two.*cells"):
            two.validate()

    def test_a_renamed_sheet_is_refused_by_name(self, tmp_path):
        source = workbook(tmp_path / "data.xlsx", {"A1": "Name"}, sheets=("One", "Two"))
        one = qd.DataSourceDefinition(source=derive(source, "XLSX").get[0])
        wb = openpyxl.load_workbook(source)
        wb["One"].title = "First"
        wb.save(source)
        with pytest.raises(ValueError, match="sheet-unknown"):
            one.validate()

    def test_a_dialect_pointed_at_the_new_name_verifies(self, tmp_path):
        source = workbook(tmp_path / "data.xlsx", {"A1": "Name"}, sheets=("One", "Two"))
        model = derive(source, "XLSX").get[0]
        wb = openpyxl.load_workbook(source)
        wb["One"].title = "First"
        wb.save(source)
        model.dialect = model.dialect.model_copy(update={"sheet": "First"})
        with pytest.warns(UserWarning, match="saved again"):
            assert qd.DataSourceDefinition(source=model).validate()

    def test_a_file_saved_again_verifies_by_fp0_with_a_warning(self, tmp_path):
        source = workbook(tmp_path / "data.xlsx", {"A1": "Name", "A2": "a"})
        datasource = derive(source, "XLSX")
        wb = openpyxl.load_workbook(source)
        wb.active.column_dimensions["A"].width = 40
        wb.save(source)
        with pytest.warns(UserWarning, match="saved again"):
            assert datasource.validate()

    def test_an_unchanged_file_verifies_without_a_warning(self, tmp_path):
        source = workbook(tmp_path / "data.xlsx", {"A1": "Name", "A2": "a"})
        datasource = derive(source, "XLSX")
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            assert datasource.validate()

    def test_a_changed_cell_fails(self, tmp_path):
        source = workbook(tmp_path / "data.xlsx", {"A1": "Name", "A2": "a"})
        datasource = derive(source, "XLSX")
        workbook(source, {"A1": "Name", "A2": "b"})
        with pytest.raises(AssertionError, match="cells"):
            datasource.validate()

    def test_a_verified_1x_checksum_gains_fp0_and_the_byte_hash(self):
        pytest.importorskip("openpyxl")
        model = CoreParser().load_json(source=DATA / "raw_E06000044_014_0.data")
        datasource = qd.DataSourceDefinition(source=model)
        assert datasource.get.get_checksum("fp0") is None
        assert datasource.validate()
        derived = derive(DATA / "raw_E06000044_014_0.XLSX", "XLSX").get
        assert derived.get_checksum("fp0") and datasource.get.get_checksum("fp0") == derived.get_checksum("fp0")
        assert datasource.get.get_checksum("bytes") == derived.get_checksum("bytes")
