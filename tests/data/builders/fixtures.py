# SPDX-License-Identifier: BSD-3-Clause
"""Write the synthetic reader fixtures, and the answer key that says what a faithful reading returns.

    python tests/data/builders/fixtures.py [target]

Target defaults to tests/data/synthetic, and the script refuses to overwrite it. Every defect is planted once, in
COLUMNS or the workbook's other sheets, and the CSVs, the workbook, and answer-key.json are all written from them.
The key is the contract: for a CSV, the exact text of every cell (an empty field is ""); for the workbook, each
cell's stored type, stored value, and number format, and the text Excel displays where that is fixed.
"""

import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import TYPE_ERROR
from workbook import save

HERE = Path(__file__).resolve().parent
TARGET = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else HERE.parent / "synthetic"

# One column per defect. "text" is the cell as a CSV carries it; "cell" (workbook only) is what is stored when it
# differs from the text, with its number format; None in "cell" means the cell is left empty.
COLUMNS = [
    ("ref_text", "a reference with leading zeros, stored as text",
     ["0000130842", "0000000001", "1234567890", "0012"], None),
    ("ref_padded", "a reference stored as a number, shown zero-padded by its number format",
     ["0000130842", "0000000001", "1234567890", "0000000012"],
     {"values": [130842, 1, 1234567890, 12], "format": "0000000000"}),
    ("amount", "currency text, including a sign after the symbol and a bare decimal",
     ["£43.40", "£-826.80", "£1,234.50", "£.00"], None),
    ("amount_mojibake", "the same with a UTF-8 £ read as Mac Roman, as Amber Valley's workbooks carry it",
     ["¬£43.40", "¬£-826.80", "¬£1,234.50", "¬£.00"], None),
    ("number_as_text", "numbers stored as text, one with a leading apostrophe-free space",
     ["12.50", "0.10", "1e3", " 7"], None),
    ("number", "numbers stored as numbers, shown to two places",
     ["12.50", "0.10", "1000.00", "7.00"], {"values": [12.5, 0.1, 1000, 7], "format": "0.00"}),
    ("blank_like", "a hard space only, spaces only, an empty cell, and a hard space inside text",
     [" ", "   ", "", "a b"], {"values": [" ", "   ", None, "a b"], "format": "General"}),
    ("padded_text", "leading and trailing spaces kept as published",
     [" lead", "trail ", " both ", "none"], None),
    ("accented", "characters every single-byte encoding here can carry",
     ["Café", "naïve", "Ætna", "Müller £"], None),
    ("date_text", "dates as text: day-first, ambiguous, year-first as pandas re-saves them, and ISO",
     ["13/04/2024", "04/01/2024", "2010-04-01 00:00:00", "2024-04-01"], None),
    ("date_mixed", "the Stockport case: date cells where Excel could read a date, text where it could not",
     ["04/01/2024", "13/04/2024", "05/02/2024", "25/12/2024"],
     {"values": [date(2024, 1, 4), "13/04/2024", date(2024, 2, 5), "25/12/2024"], "format": "dd/mm/yyyy"}),
    ("date_locale", "a date cell in Excel's built-in short date format, whose display follows the machine's locale",
     [None, None, None, None], {"values": [date(2024, 1, 4), date(2024, 12, 25), datetime(2024, 4, 1, 9, 30), None],
                                "format": "mm-dd-yy"}),
    ("boolean_like", "boolean cells and boolean-looking text",
     ["TRUE", "false", "Y", "FALSE"], {"values": [True, "false", "Y", False], "format": "General"}),
]
ROWS = 4
ENCODINGS = [("utf-8", "utf8"), ("utf-8-sig", "utf8-bom"), ("utf-16", "utf16-bom"), ("cp1252", "windows-1252"),
             ("latin-1", "latin-1"), ("mac_roman", "mac-roman")]
UNICODE_ONLY = [
    ("zero_width", "a zero-width space and a byte-order mark inside text", ["abc​", "﻿x"]),
    ("control", "a bell character and a tab inside a quoted field", ["x\u0007y", "a\tb"]),
    ("composed", "é composed (NFC) and decomposed (NFD)", ["café", "café"]),
    ("fullwidth", "full-width digits and a Unicode minus", ["１２３", "−5"]),
]


def csv_text(header, rows, eol="\n"):
    def field(value):
        value = "" if value is None else value
        if any(c in value for c in ',"\n\r\t') or value != value.strip():
            return '"' + value.replace('"', '""') + '"'
        return value
    return eol.join(",".join(field(v) for v in row) for row in [header, *rows]) + eol


def main():
    assert not TARGET.exists(), f"{TARGET} exists; fixtures are never overwritten"
    TARGET.mkdir(parents=True)
    key = {"generated": datetime.now().isoformat(timespec="seconds"), "columns": {n: d for n, d, _, _ in COLUMNS},
           "csv": {}, "workbook": {}}

    # CSV: the text of every column, in each encoding, LF endings; one UTF-8 file with CRLF, one with both.
    csv_columns = [c for c in COLUMNS if c[2][0] is not None]
    header = [c[0] for c in csv_columns]
    rows = [[c[2][i] for c in csv_columns] for i in range(ROWS)]
    for codec, label in ENCODINGS:
        name = f"fixture-{label}.csv"
        (TARGET / name).write_bytes(csv_text(header, rows).encode(codec))
        key["csv"][name] = {"encoding": codec, "header": header, "rows": rows}
    (TARGET / "fixture-utf8-crlf.csv").write_bytes(csv_text(header, rows, "\r\n").encode("utf-8"))
    key["csv"]["fixture-utf8-crlf.csv"] = {"encoding": "utf-8", "eol": "CRLF", "header": header, "rows": rows}
    mixed = csv_text(header, rows[:2]) + csv_text(header, rows[2:], "\r\n").split("\r\n", 1)[1]
    (TARGET / "fixture-utf8-mixed-eol.csv").write_bytes(mixed.encode("utf-8"))
    key["csv"]["fixture-utf8-mixed-eol.csv"] = {"encoding": "utf-8", "eol": "LF then CRLF", "header": header, "rows": rows}
    u_header = [c[0] for c in UNICODE_ONLY]
    u_rows = [[c[2][i] for c in UNICODE_ONLY] for i in range(2)]
    (TARGET / "fixture-unicode-utf8.csv").write_bytes(csv_text(u_header, u_rows).encode("utf-8"))
    key["csv"]["fixture-unicode-utf8.csv"] = {"encoding": "utf-8", "header": u_header, "rows": u_rows,
                                              "columns": {n: d for n, d, _ in UNICODE_ONLY}}

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    cells = {}
    for j, (name, _, text, cell) in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=j, value=name)
        for i in range(ROWS):
            value = cell["values"][i] if cell else text[i]
            c = ws.cell(row=i + 2, column=j, value=value)
            if cell:
                c.number_format = cell["format"]
            cells[c.coordinate] = {"column": name, "stored_type": c.data_type, "stored": value,
                                   "format": c.number_format, "text": text[i]}
    err = ws.cell(row=ROWS + 2, column=1, value="#N/A")
    err.data_type = TYPE_ERROR
    cells[err.coordinate] = {"column": "ref_text", "stored_type": "e", "stored": "#N/A", "format": "General", "text": "#N/A",
                             "note": "an error cell, below the data"}
    f = ws.cell(row=ROWS + 2, column=6, value="=F2*2")
    cells[f.coordinate] = {"column": "number", "stored_type": "f", "stored": "=F2*2", "format": "General", "text": "25",
                           "note": "a formula; written here without a cached value, which only Excel adds when it saves"}
    key["workbook"]["Data"] = {"header_row": 1, "cells": cells,
                               "note": "date_locale's display depends on the locale (built-in format 14): en-GB 04/01/2024, en-US 1/4/2024"}

    titled = wb.create_sheet("Titled")
    titled["A1"] = "Reliefs key"
    for r, row in enumerate([["Code", "Description"], ["EPCH", "Empty Property Rate Charitable"], ["MAN", "Mandatory"]], start=3):
        for c, value in enumerate(row, start=1):
            titled.cell(row=r, column=c, value=value)
    key["workbook"]["Titled"] = {"layout": "title in A1, row 2 blank, header in row 3, two data rows",
                                 "rows": [["Reliefs key", None], [None, None], ["Code", "Description"],
                                          ["EPCH", "Empty Property Rate Charitable"], ["MAN", "Mandatory"]]}
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"], hidden["A2"] = "secret", "kept"
    key["workbook"]["Hidden"] = {"state": "hidden", "rows": [["secret"], ["kept"]]}
    wb.create_sheet("Empty")
    key["workbook"]["Empty"] = {"rows": []}
    save(wb, TARGET / "fixture.xlsx")
    key["workbook_file"] = "fixture.xlsx"

    (TARGET / "answer-key.json").write_text(json.dumps(key, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    print(sorted(p.name for p in TARGET.iterdir()))


if __name__ == "__main__":
    main()
